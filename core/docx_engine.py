"""DOCX Replacement Engine and Post-Processing Verifier.

Coordinates the replacement across all document structures (body, tables, headers, footers),
enforces source immutability, manages output naming, and verifies output validity.
"""

from pathlib import Path
from typing import Optional, Set
import docx
from docx.text.paragraph import Paragraph

from core.models import (
    ConflictResolution,
    FileProcessStatus,
    IdentityConfig,
    ProcessingResult,
    ReplacementMode,
    TargetMatchCount,
)
from core.run_replacer import replace_text_in_paragraph
from core.docx_validator import scan_document_for_matches


def resolve_unique_filename(destination_dir: Path, filename: str) -> Path:
    """Generates a non-conflicting filename by appending _rev1, _rev2, etc."""
    base_path = destination_dir / filename
    if not base_path.exists():
        return base_path

    stem = Path(filename).stem
    suffix = Path(filename).suffix
    counter = 1
    while True:
        candidate = destination_dir / f"{stem}_rev{counter}{suffix}"
        if not candidate.exists():
            return candidate
        counter += 1


def apply_replacements_to_paragraphs(
    paragraphs: list[Paragraph],
    old_target: str,
    new_target: str,
) -> int:
    """Applies replacement across a list of paragraphs using run-aware replacer."""
    if not old_target or not new_target:
        return 0

    count = 0
    for p in paragraphs:
        count += replace_text_in_paragraph(p, old_target, new_target)
    return count


class DocxEngine:
    """Orchestrates document replacement and verification."""

    @staticmethod
    def process_file(
        source_file: Path,
        output_dir: Path,
        old_identity: IdentityConfig,
        new_identity: IdentityConfig,
        mode: ReplacementMode = ReplacementMode.BOTH,
        conflict_resolution: ConflictResolution = ConflictResolution.AUTO_RENAME,
    ) -> ProcessingResult:
        """Processes a single document: replaces identity, saves new file, and verifies.
        
        Strictly guarantees source file immutability (BR-005).
        """
        src = Path(source_file).resolve()
        out_dir = Path(output_dir).resolve()

        # 1. Basic validation
        if not src.exists():
            return ProcessingResult(
                source_file=src,
                status=FileProcessStatus.FAILED,
                status_message=f"File sumber {src.name} tidak ditemukan.",
            )

        if src.suffix.lower() != ".docx":
            return ProcessingResult(
                source_file=src,
                status=FileProcessStatus.FAILED,
                status_message=f"File {src.name} bukan format .docx.",
            )

        # 2. Determine target values based on mode
        replace_name = mode in (ReplacementMode.NAME_ONLY, ReplacementMode.BOTH)
        replace_nim = mode in (ReplacementMode.NIM_ONLY, ReplacementMode.BOTH)

        target_old_name = old_identity.name if replace_name else ""
        target_new_name = new_identity.name if replace_name else ""

        # Keep NIM strictly string
        target_old_nim = str(old_identity.nim) if replace_nim else ""
        target_new_nim = str(new_identity.nim) if replace_nim else ""

        # 3. Load document into memory
        try:
            doc = docx.Document(src)
        except Exception as e:
            return ProcessingResult(
                source_file=src,
                status=FileProcessStatus.FAILED,
                status_message=f"Gagal membuka dokumen {src.name}: {str(e)}",
            )

        # 4. Perform replacements
        total_name_replaced = 0
        total_nim_replaced = 0

        # Helper to process any paragraph list
        def process_p_list(plist: list[Paragraph]) -> None:
            nonlocal total_name_replaced, total_nim_replaced
            if target_old_name and target_new_name:
                total_name_replaced += apply_replacements_to_paragraphs(
                    plist, target_old_name, target_new_name
                )
            if target_old_nim and target_new_nim:
                total_nim_replaced += apply_replacements_to_paragraphs(
                    plist, target_old_nim, target_new_nim
                )

        # 4a. Process Body Paragraphs
        process_p_list(doc.paragraphs)

        # 4b. Process Tables
        for table in doc.tables:
            for row in table.rows:
                for cell in row.cells:
                    process_p_list(cell.paragraphs)

        # 4c. Process Headers & Footers
        seen_parts: Set[int] = set()
        for section in doc.sections:
            for hf_part in (
                section.header,
                section.footer,
                section.first_page_header,
                section.first_page_footer,
                section.even_page_header,
                section.even_page_footer,
            ):
                if hf_part is None:
                    continue
                part_id = id(hf_part._element)
                if part_id in seen_parts:
                    continue
                seen_parts.add(part_id)

                process_p_list(hf_part.paragraphs)
                for table in hf_part.tables:
                    for row in table.rows:
                        for cell in row.cells:
                            process_p_list(cell.paragraphs)

        matches_replaced = TargetMatchCount(
            name_matches=total_name_replaced, nim_matches=total_nim_replaced
        )

        # 5. Check if any replacement happened (BR-006)
        total_replaced = 0
        if mode == ReplacementMode.NAME_ONLY:
            total_replaced = total_name_replaced
        elif mode == ReplacementMode.NIM_ONLY:
            total_replaced = total_nim_replaced
        else:
            total_replaced = total_name_replaced + total_nim_replaced

        if total_replaced == 0:
            return ProcessingResult(
                source_file=src,
                status=FileProcessStatus.NO_TARGET,
                status_message=f"Identitas lama tidak ditemukan pada file {src.name}.",
                matches_replaced=matches_replaced,
            )

        # 6. Determine destination path & resolve conflicts
        out_dir.mkdir(parents=True, exist_ok=True)
        dest_file = out_dir / src.name

        # Enforce source file immutability: never allow dest_file to point to src!
        if dest_file.resolve() == src.resolve():
            dest_file = resolve_unique_filename(out_dir, f"{src.stem}_output{src.suffix}")
        elif dest_file.exists():
            if conflict_resolution == ConflictResolution.AUTO_RENAME:
                dest_file = resolve_unique_filename(out_dir, src.name)
            elif conflict_resolution == ConflictResolution.SKIP:
                return ProcessingResult(
                    source_file=src,
                    status=FileProcessStatus.CANCELLED,
                    status_message=f"File {dest_file.name} dilewati karena sudah ada di folder keluaran.",
                )
            # if OVERWRITE, proceed with dest_file as is

        # 7. Save to destination
        try:
            doc.save(dest_file)
        except Exception as e:
            return ProcessingResult(
                source_file=src,
                status=FileProcessStatus.FAILED,
                status_message=f"Gagal menyimpan file hasil {dest_file.name}: {str(e)}",
            )

        # 8. Post-replacement Verification (FR-014, AC-011)
        try:
            verified_doc = docx.Document(dest_file)
            search_old_name = target_old_name
            search_old_nim = target_old_nim
            post_counts = scan_document_for_matches(
                verified_doc, search_old_name, search_old_nim
            )

            # Old target should no longer appear
            old_remnants = 0
            if replace_name:
                old_remnants += post_counts.name_matches
            if replace_nim:
                old_remnants += post_counts.nim_matches

            # Scan if new targets appear
            new_counts = scan_document_for_matches(
                verified_doc, target_new_name, target_new_nim
            )
            new_present = 0
            if replace_name:
                new_present += new_counts.name_matches
            if replace_nim:
                new_present += new_counts.nim_matches

            if old_remnants == 0 and new_present > 0:
                return ProcessingResult(
                    source_file=src,
                    output_file=dest_file,
                    status=FileProcessStatus.SUCCESS,
                    status_message=f"Berhasil mengganti {total_replaced} identitas target.",
                    matches_replaced=matches_replaced,
                    verified=True,
                )
            else:
                return ProcessingResult(
                    source_file=src,
                    output_file=dest_file,
                    status=FileProcessStatus.NEEDS_REVIEW,
                    status_message="Hasil pada file tidak dapat diverifikasi sepenuhnya. Periksa secara manual.",
                    matches_replaced=matches_replaced,
                    verified=False,
                    warning="Verifikasi pasca-proses mendeteksi sisa identitas lama atau ketidakcocokan identitas baru.",
                )
        except Exception as e:
            return ProcessingResult(
                source_file=src,
                output_file=dest_file,
                status=FileProcessStatus.NEEDS_REVIEW,
                status_message=f"File hasil tersimpan tetapi gagal diverifikasi: {str(e)}",
                matches_replaced=matches_replaced,
                verified=False,
            )
