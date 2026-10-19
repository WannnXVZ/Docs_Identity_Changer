"""Test Document Generator for Word Identity Changer.

Generates realistic .docx test documents covering all PRD Section O scenarios:
1. Standard paragraphs
2. Split runs across different styles (bold, italic, colors)
3. Table cells
4. Headers and Footers
5. Leading zero NIMs
6. Substring collision prevention (e.g. 12345 vs 9123456)
7. Document with no target identity
8. Multiple occurrences
9. Corrupted file
"""

from pathlib import Path
import docx
from docx.shared import Pt, RGBColor


def generate_all_test_documents(output_dir: Path) -> dict[str, Path]:
    """Generates all 9 test documents in output_dir and returns mapping of names to paths."""
    output_dir.mkdir(parents=True, exist_ok=True)
    generated_files = {}

    OLD_NAME = "Budi Santoso"
    OLD_NIM = "0701234567"

    # 1. Simple Paragraph
    doc1 = docx.Document()
    doc1.add_heading("Laporan Praktikum Fisika Dasar", level=1)
    doc1.add_paragraph(f"Nama Mahasiswa: {OLD_NAME}")
    doc1.add_paragraph(f"Nomor Induk Mahasiswa: {OLD_NIM}")
    doc1.add_paragraph("Praktikum ini membahas tentang hukum Archimedes pada fluida cair.")
    p1 = output_dir / "test_doc1_simple_paragraph.docx"
    doc1.save(p1)
    generated_files["doc1"] = p1

    # 2. Split Runs (Crucial for testing Run-Aware Replacer!)
    doc2 = docx.Document()
    doc2.add_heading("Tugas Pemrograman Web", level=1)
    p = doc2.add_paragraph()
    p.add_run("Disusun oleh: ")
    r_name1 = p.add_run("Budi ")
    r_name1.bold = True
    r_name2 = p.add_run("Santoso")
    r_name2.italic = True

    p_nim = doc2.add_paragraph()
    p_nim.add_run("NIM: ")
    r_nim1 = p_nim.add_run("07012")
    r_nim1.font.color.rgb = RGBColor(0, 51, 102)
    r_nim2 = p_nim.add_run("34567")
    r_nim2.bold = True
    doc2.add_paragraph("Laporan tugas besar membuat aplikasi e-commerce.")
    p2 = output_dir / "test_doc2_split_runs.docx"
    doc2.save(p2)
    generated_files["doc2"] = p2

    # 3. Identity in Table
    doc3 = docx.Document()
    doc3.add_heading("Form Penilaian Tugas Akhir", level=1)
    table = doc3.add_table(rows=3, cols=2)
    table.style = "Table Grid"
    table.cell(0, 0).paragraphs[0].text = "Informasi"
    table.cell(0, 1).paragraphs[0].text = "Keterangan"
    table.cell(1, 0).paragraphs[0].text = "Nama Mahasiswa"
    table.cell(1, 1).paragraphs[0].text = OLD_NAME
    table.cell(2, 0).paragraphs[0].text = "NIM"
    table.cell(2, 1).paragraphs[0].text = OLD_NIM
    p3 = output_dir / "test_doc3_in_table.docx"
    doc3.save(p3)
    generated_files["doc3"] = p3

    # 4. In Header and Footer
    doc4 = docx.Document()
    section = doc4.sections[0]
    header_p = section.header.paragraphs[0]
    header_p.text = f"Dokumen Resmi - {OLD_NAME} ({OLD_NIM})"
    footer_p = section.footer.paragraphs[0]
    footer_p.text = f"Halaman 1 - Pemilik: {OLD_NAME}"
    doc4.add_heading("Dokumen dengan Header dan Footer", level=1)
    doc4.add_paragraph("Isi konten dokumen utama...")
    p4 = output_dir / "test_doc4_header_footer.docx"
    doc4.save(p4)
    generated_files["doc4"] = p4

    # 5. Leading Zero NIM
    doc5 = docx.Document()
    doc5.add_heading("Verifikasi NIM Berawalan Nol", level=1)
    doc5.add_paragraph(f"Nama: {OLD_NAME}")
    doc5.add_paragraph(f"NIM: {OLD_NIM}")
    p5 = output_dir / "test_doc5_leading_zero.docx"
    doc5.save(p5)
    generated_files["doc5"] = p5

    # 6. Substring Collision Prevention (BR-009)
    # Testing that '12345' does NOT match inside '9123456'
    doc6 = docx.Document()
    doc6.add_heading("Pencegahan Substring Match", level=1)
    doc6.add_paragraph(f"Mahasiswa Utama: {OLD_NAME} - NIM: 12345")
    doc6.add_paragraph("Rekan Lain: Anton - NIM: 9123456 (Tidak boleh berubah!)")
    doc6.add_paragraph("Kode Referensi Barang: REF12345XYZ (Tidak boleh berubah!)")
    p6 = output_dir / "test_doc6_substring_safety.docx"
    doc6.save(p6)
    generated_files["doc6"] = p6

    # 7. Document without Target Identity
    doc7 = docx.Document()
    doc7.add_heading("Laporan Mahasiswa Lain", level=1)
    doc7.add_paragraph("Nama: Siti Aminah")
    doc7.add_paragraph("NIM: 0998877665")
    p7 = output_dir / "test_doc7_no_target.docx"
    doc7.save(p7)
    generated_files["doc7"] = p7

    # 8. Multiple Occurrences
    doc8 = docx.Document()
    doc8.add_heading("Dokumen dengan Banyak Kemunculan", level=1)
    doc8.add_paragraph(f"Cover: {OLD_NAME} - {OLD_NIM}")
    doc8.add_paragraph(f"Kata Pengantar: Ditulis oleh {OLD_NAME}.")
    doc8.add_paragraph(f"Lembar Pengesahan: Mahasiswa bersangkutan {OLD_NAME} ({OLD_NIM}).")
    t8 = doc8.add_table(rows=2, cols=2)
    t8.cell(0, 0).paragraphs[0].text = "Nama"
    t8.cell(0, 1).paragraphs[0].text = OLD_NAME
    t8.cell(1, 0).paragraphs[0].text = "NIM"
    t8.cell(1, 1).paragraphs[0].text = OLD_NIM
    p8 = output_dir / "test_doc8_multiple.docx"
    doc8.save(p8)
    generated_files["doc8"] = p8

    # 9. Corrupt File
    p9 = output_dir / "test_doc9_corrupted.docx"
    p9.write_bytes(b"PK\x03\x04This is a completely corrupted zip content not a real docx")
    generated_files["doc9"] = p9

    return generated_files


if __name__ == "__main__":
    test_dir = Path(__file__).parent / "sample_docs"
    created = generate_all_test_documents(test_dir)
    print(f"Successfully generated {len(created)} test documents in {test_dir}:")
    for k, v in created.items():
        print(f"  [{k}] -> {v.name}")
