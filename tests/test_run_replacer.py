"""Unit tests for Run-Aware Text Replacer Engine."""

import unittest
import docx
from core.run_replacer import (
    build_boundary_regex,
    count_matches_in_paragraph,
    replace_text_in_paragraph,
)


class TestRunReplacer(unittest.TestCase):
    def test_boundary_regex_substring_safety(self):
        """Tests that boundary regex does not match substrings (BR-009)."""
        pattern = build_boundary_regex("12345")
        
        # Exact match
        self.assertEqual(len(pattern.findall("NIM: 12345")), 1)
        self.assertEqual(len(pattern.findall("12345")), 1)
        self.assertEqual(len(pattern.findall("NIM 12345 dan nama")), 1)

        # Substring inside another number - MUST NOT MATCH!
        self.assertEqual(len(pattern.findall("NIM: 9123456")), 0)
        self.assertEqual(len(pattern.findall("REF12345XYZ")), 0)
        self.assertEqual(len(pattern.findall("123450")), 0)

    def test_single_run_replacement(self):
        """Tests replacement within a single run."""
        doc = docx.Document()
        p = doc.add_paragraph("Nama mahasiswa adalah Budi Santoso di kelas A.")
        
        count = replace_text_in_paragraph(p, "Budi Santoso", "Andi Wijaya")
        self.assertEqual(count, 1)
        self.assertEqual(p.text, "Nama mahasiswa adalah Andi Wijaya di kelas A.")

    def test_split_runs_replacement(self):
        """Tests replacement across multiple runs with different formatting."""
        doc = docx.Document()
        p = doc.add_paragraph()
        r1 = p.add_run("Nama: ")
        r2 = p.add_run("Budi ")
        r2.bold = True
        r3 = p.add_run("Santoso")
        r3.italic = True
        r4 = p.add_run(" (Selesai)")

        # Verify initial text
        self.assertEqual(p.text, "Nama: Budi Santoso (Selesai)")
        
        # Count matches across runs
        matches = count_matches_in_paragraph(p, "Budi Santoso")
        self.assertEqual(matches, 1)

        # Replace across runs
        count = replace_text_in_paragraph(p, "Budi Santoso", "Dewi Sartika")
        self.assertEqual(count, 1)
        self.assertEqual(p.text, "Nama: Dewi Sartika (Selesai)")
        
        # Verify first run retained its bold styling
        self.assertTrue(r2.bold)
        # Verify text was cleaned from r3
        self.assertEqual(r3.text, "")

    def test_leading_zero_preservation(self):
        """Tests NIM with leading zero is preserved as exact string (BR-004)."""
        doc = docx.Document()
        p = doc.add_paragraph("NIM Lama: 0701234567")

        count = replace_text_in_paragraph(p, "0701234567", "0809876543")
        self.assertEqual(count, 1)
        self.assertEqual(p.text, "NIM Lama: 0809876543")
        self.assertTrue(p.text.endswith("0809876543"))


if __name__ == "__main__":
    unittest.main()
